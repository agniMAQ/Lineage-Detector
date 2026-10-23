#!/usr/bin/env python3

from __future__ import annotations

from collections import Counter
from pathlib import Path
import re

import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter


WORKSPACE = Path(r"c:\Repos\WK LR US Artifacts")
LINEAGE_DIR = WORKSPACE / "Lineage Viewer" / "lineage-output"
ANALYSIS_DIR = WORKSPACE / "Lineage Viewer" / "analysis"
OUTPUT_FILE = ANALYSIS_DIR / "lineage_detailed_analysis.xlsx"


CONSOLIDATED_LINEAGE_CSV = LINEAGE_DIR / "lineage_metadata_bi_reports_only.csv"


PLATFORM_MAP = {
    "tgoogleanalyticsinput": "GOOGLE_ANALYTICS",
    "tgoogleanalytics4input": "GOOGLE_ANALYTICS",
    "tsalesforceinput": "SALESFORCE",
    "tsalesforcerecurrenceinput": "SALESFORCE",
    "tgainsightinput": "GAINSIGHT",
    "tredshiftinput": "REDSHIFT",
    "tdbinput": "REDSHIFT",
    "tjdbcinput": "REDSHIFT",
    "toracleinput": "ORACLE",
    "tmysqlinput": "MYSQL",
    "tmssqlinput": "MSSQL",
    "tfileinputdelimited": "FILE",
    "tfileinputjson": "FILE",
    "tfileinputexcel": "FILE",
    "tfileinputraw": "FILE",
    "tftpget": "FILE",
    "ts3get": "S3",
    "thttpinput": "API",
    "trestclient": "API",
    "thttprequest": "API",
    "trest": "API",
}


def read_csv(path: Path) -> pd.DataFrame:
    if not path.exists():
        return pd.DataFrame()
    return pd.read_csv(path, dtype=str).fillna("")


def normalize_columns(df: pd.DataFrame) -> pd.DataFrame:
    if df.empty:
        return df
    out = df.copy()
    out.columns = [str(c).strip() for c in out.columns]
    return out


def sheet_name(name: str) -> str:
    return name[:31]


def write_sheet(writer: pd.ExcelWriter, name: str, df: pd.DataFrame) -> None:
    df = normalize_columns(df)
    df.to_excel(writer, sheet_name=sheet_name(name), index=False)


def apply_formatting(workbook_path: Path) -> None:
    wb = load_workbook(workbook_path)
    header_fill = PatternFill("solid", fgColor="1F4E78")
    header_font = Font(color="FFFFFF", bold=True)
    thin_alignment = Alignment(vertical="top", wrap_text=True)

    for ws in wb.worksheets:
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions
        ws.sheet_view.showGridLines = False

        for cell in ws[1]:
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = thin_alignment

        for row in ws.iter_rows(min_row=2):
            for cell in row:
                cell.alignment = thin_alignment

        widths = {}
        for row in ws.iter_rows(values_only=True):
            for idx, value in enumerate(row, start=1):
                text = "" if value is None else str(value)
                widths[idx] = max(widths.get(idx, 0), min(len(text), 80))
        for idx, width in widths.items():
            ws.column_dimensions[get_column_letter(idx)].width = max(12, min(width + 2, 40))

    wb.save(workbook_path)


def is_valid_name(value: str) -> bool:
    text = str(value).strip()
    return text != "" and text.upper() != "NULL"


def normalize_identifier(value: str) -> str:
    """Normalize schema/object identifiers and remove stray quote wrappers."""
    text = str(value).strip().lower()
    # Remove common quote/bracket wrappers produced by CSV/SQL formatting noise.
    text = text.replace('"', "").replace("'", "")
    text = text.replace("[", "").replace("]", "")
    return text.strip()


def normalize_key(schema: str, obj: str) -> tuple[str, str]:
    schema_norm = normalize_identifier(schema)
    obj_norm = normalize_identifier(obj)

    # Some upstream rows include malformed split values like:
    #   schema="foo (bi" and object="foo) (bi)"
    # Recover canonical "bi.foo" when present anywhere in the combined token.
    combined = f"{schema_norm}.{obj_norm}"
    m = re.search(r"\(([a-z0-9_]+)\.([a-z0-9_]+)\)", combined)
    if m:
        return m.group(1), m.group(2)

    # Also accept direct schema.table if one side accidentally contains both parts.
    if "." in schema_norm and obj_norm:
        left, right = schema_norm.split(".", 1)
        if right == obj_norm:
            return left, right

    return schema_norm, obj_norm


def used_redshift_objects(df: pd.DataFrame) -> pd.DataFrame:
    if df.empty:
        return pd.DataFrame()

    src = df[df["Source_Platform"].astype(str).str.upper() == "REDSHIFT"].copy()
    src = src[src["Source_Schema"].map(is_valid_name) & src["Source_Object"].map(is_valid_name)]

    tgt = df[df["Target_Platform"].astype(str).str.upper() == "REDSHIFT"].copy()
    tgt = tgt[tgt["Target_Schema"].map(is_valid_name) & tgt["Target_Object"].map(is_valid_name)]

    all_keys = set(
        [normalize_key(r["Source_Schema"], r["Source_Object"]) for _, r in src.iterrows()]
        + [normalize_key(r["Target_Schema"], r["Target_Object"]) for _, r in tgt.iterrows()]
    )

    rows = []
    for schema_name, object_name in sorted(all_keys):
        src_mask = (
            src["Source_Schema"].astype(str).str.lower().eq(schema_name)
            & src["Source_Object"].astype(str).str.lower().eq(object_name)
        )
        tgt_mask = (
            tgt["Target_Schema"].astype(str).str.lower().eq(schema_name)
            & tgt["Target_Object"].astype(str).str.lower().eq(object_name)
        )

        downstream_platforms = sorted(set(src.loc[src_mask, "Target_Platform"].astype(str)) - {""})
        upstream_platforms = sorted(set(tgt.loc[tgt_mask, "Source_Platform"].astype(str)) - {""})

        rows.append(
            {
                "object_name": f"{schema_name}.{object_name}",
                "schema_name": schema_name,
                "table_name": object_name,
                "used_as_source_rows": int(src_mask.sum()),
                "used_as_target_rows": int(tgt_mask.sum()),
                "used_as_source": "Yes" if int(src_mask.sum()) > 0 else "No",
                "used_as_target": "Yes" if int(tgt_mask.sum()) > 0 else "No",
                "downstream_platforms": "; ".join(downstream_platforms),
                "upstream_platforms": "; ".join(upstream_platforms),
            }
        )

    return pd.DataFrame(rows)


def used_transform_files(df: pd.DataFrame, tool_name: str, label: str) -> pd.DataFrame:
    if df.empty:
        return pd.DataFrame()

    out = df[df["Transform_Tool"].astype(str).str.upper() == tool_name.upper()].copy()
    if out.empty:
        return pd.DataFrame()

    out["transform_path_norm"] = out["Transform_Path"].astype(str).str.strip()
    out["transform_name_norm"] = out["Transform_Name"].astype(str).str.strip()

    grouped = (
        out.groupby(["transform_path_norm", "transform_name_norm"], dropna=False)
        .agg(
            lineage_rows=("Transform_Path", "size"),
            unique_source_objects=("Source_Object", "nunique"),
            unique_target_objects=("Target_Object", "nunique"),
            source_platforms=("Source_Platform", lambda s: "; ".join(sorted(set(v for v in s.astype(str) if v)))),
            target_platforms=("Target_Platform", lambda s: "; ".join(sorted(set(v for v in s.astype(str) if v)))),
        )
        .reset_index()
    )

    grouped = grouped.rename(
        columns={
            "transform_path_norm": f"{label}_path",
            "transform_name_norm": f"{label}_name",
        }
    )
    return grouped.sort_values(by=["lineage_rows", f"{label}_path"], ascending=[False, True]).reset_index(drop=True)


def redshift_no_source_usage(df: pd.DataFrame) -> pd.DataFrame:
    if df.empty:
        return pd.DataFrame()

    src = df[df["Source_Platform"].astype(str).str.upper() == "REDSHIFT"].copy()
    src = src[src["Source_Schema"].map(is_valid_name) & src["Source_Object"].map(is_valid_name)]

    tgt = df[df["Target_Platform"].astype(str).str.upper() == "REDSHIFT"].copy()
    tgt = tgt[tgt["Target_Schema"].map(is_valid_name) & tgt["Target_Object"].map(is_valid_name)]

    source_keys = set(normalize_key(r["Source_Schema"], r["Source_Object"]) for _, r in src.iterrows())
    target_keys = set(normalize_key(r["Target_Schema"], r["Target_Object"]) for _, r in tgt.iterrows())

    # Objects that are used as Redshift sources but never appear as Redshift targets.
    no_source_keys = sorted(source_keys - target_keys)

    rows = []
    for schema_name, object_name in no_source_keys:
        mask = (
            src["Source_Schema"].astype(str).str.lower().eq(schema_name)
            & src["Source_Object"].astype(str).str.lower().eq(object_name)
        )
        usage_rows = src.loc[mask].copy()

        target_platforms = sorted(set(usage_rows["Target_Platform"].astype(str)) - {""})
        target_platforms_upper = {p.upper() for p in target_platforms}

        used_in_power_bi = any(p.startswith("POWER_BI") for p in target_platforms_upper)
        used_in_tableau = "TABLEAU" in target_platforms_upper
        used_in_talend = "TALEND" in set(usage_rows["Transform_Tool"].astype(str).str.upper())

        rows.append(
            {
                "object_name": f"{schema_name}.{object_name}",
                "schema_name": schema_name,
                "table_name": object_name,
                "downstream_rows": len(usage_rows),
                "downstream_platforms": "; ".join(target_platforms),
                "downstream_transform_tools": "; ".join(sorted(set(v for v in usage_rows["Transform_Tool"].astype(str) if v))),
                "used_in_talend": "Yes" if used_in_talend else "No",
                "used_in_power_bi": "Yes" if used_in_power_bi else "No",
                "used_in_tableau": "Yes" if used_in_tableau else "No",
                "example_transform_paths": "; ".join(sorted(set(v for v in usage_rows["Transform_Path"].astype(str) if v))[:10]),
            }
        )

    return pd.DataFrame(rows)


def main() -> None:
    ANALYSIS_DIR.mkdir(parents=True, exist_ok=True)

    lineage_df = read_csv(CONSOLIDATED_LINEAGE_CSV)

    redshift_objects_df = used_redshift_objects(lineage_df)
    python_files_df = used_transform_files(lineage_df, tool_name="PYTHON", label="python_file")
    talend_pipelines_df = used_transform_files(lineage_df, tool_name="TALEND", label="talend_pipeline")
    redshift_no_source_df = redshift_no_source_usage(lineage_df)

    overview_df = pd.DataFrame([
        {"dataset": "Consolidated lineage", "rows": int(len(lineage_df)), "status": "available" if not lineage_df.empty else "empty or missing"},
        {"dataset": "Used Redshift objects", "rows": int(len(redshift_objects_df)), "status": "generated"},
        {"dataset": "Used Python files", "rows": int(len(python_files_df)), "status": "generated"},
        {"dataset": "Used Talend pipelines", "rows": int(len(talend_pipelines_df)), "status": "generated"},
        {"dataset": "Redshift no-source usage", "rows": int(len(redshift_no_source_df)), "status": "generated"},
    ])

    with pd.ExcelWriter(OUTPUT_FILE, engine="openpyxl") as writer:
        write_sheet(writer, "Overview", overview_df)
        write_sheet(writer, "Used_Redshift_Objects", redshift_objects_df)
        write_sheet(writer, "Used_Python_Files", python_files_df)
        write_sheet(writer, "Used_Talend_Pipelines", talend_pipelines_df)
        write_sheet(writer, "Redshift_No_Source_Usage", redshift_no_source_df)

    apply_formatting(OUTPUT_FILE)

    print(f"Wrote {OUTPUT_FILE}")
    print(f"Consolidated lineage rows: {len(lineage_df)}")
    print(f"Used Redshift objects: {len(redshift_objects_df)}")
    print(f"Used Python files: {len(python_files_df)}")
    print(f"Used Talend pipelines: {len(talend_pipelines_df)}")
    print(f"Redshift no-source usage rows: {len(redshift_no_source_df)}")


if __name__ == "__main__":
    main()